"""
Importação de DUIMP → Planilha Omie — Front Streamlit (Sillion)
Fluxo:
1. Usuário informa o número da DUIMP (+ versão e email).
2. O app chama o webhook do N8N, que consulta a DUIMP no Portal Único,
   resolve/cria os produtos no Omie, classifica CFOP/categoria (IA) e
   devolve os dados estruturados da planilha de Importação de Mercadoria.
3. O app preenche o template oficial (openpyxl) e oferece o download.

Arquitetura:
- app.py        → lógica Python (envio, montagem da planilha, download)
- styles/       → CSS
- templates/    → HTML estrutural (header, hero, footer)
- template_importacao_mercadoria.xlsx → modelo oficial do Omie
"""

import re
import io
from datetime import datetime
from pathlib import Path

import requests
import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

# ============================================================
# Caminhos
# ============================================================
BASE_DIR = Path(__file__).parent
CSS_PATH = BASE_DIR / "styles" / "main.css"
TEMPLATES_DIR = BASE_DIR / "templates"
STATIC_DIR = BASE_DIR / "static"
TEMPLATE_XLSX = BASE_DIR / "template_importacao_mercadoria.xlsx"
SHEET_PEDIDO = "Omie_Importacao_Mercadoria"
LINHA_CABECALHO = 7       # linha de dados do cabeçalho do pedido
LINHA_PRIMEIRO_ITEM = 17  # primeira linha de dados dos itens

LOGO_EXTERNO = "https://www.sillion.com.br/wp-content/themes/sillion/images/logo-black-tm.svg"
LOGO_LOCAL_FILE = STATIC_DIR / "logo-sillion.svg"


def resolver_logo_url() -> str:
    if LOGO_LOCAL_FILE.exists():
        return "app/static/logo-sillion.svg"
    return LOGO_EXTERNO


st.set_page_config(
    page_title="Sillion · DUIMP → Importação Omie",
    page_icon="https://www.sillion.com.br/wp-content/themes/sillion/images/logo-white-tm.svg",
    layout="centered",
    initial_sidebar_state="collapsed",
)

DOMINIO_PERMITIDO = "sillion.com.br"
EMAIL_REGEX = re.compile(rf"^[A-Za-z0-9._%+\-]+@{re.escape(DOMINIO_PERMITIDO)}$", re.IGNORECASE)
TIMEOUT_REQ = 180  # segundos (o fluxo consulta a DUIMP + IA + Omie)
COL_RE = re.compile(r"^[A-Z]{1,3}$")
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ============================================================
# Helpers de renderização (templates + CSS)
# ============================================================
def render_template(nome: str, **variaveis) -> str:
    html = (TEMPLATES_DIR / f"{nome}.html").read_text(encoding="utf-8")
    for chave, valor in variaveis.items():
        html = html.replace(f"{{{{{chave}}}}}", str(valor))
    return html


def inject(html: str) -> None:
    st.markdown(html, unsafe_allow_html=True)


def carregar_css(caminho: Path) -> None:
    try:
        inject(f"<style>{caminho.read_text(encoding='utf-8')}</style>")
    except FileNotFoundError:
        st.warning(f"Arquivo de estilos não encontrado: {caminho}")


inject(render_template("meta"))
carregar_css(CSS_PATH)

try:
    WEBHOOK_URL = st.secrets["N8N_WEBHOOK_URL"]
except (KeyError, FileNotFoundError):
    WEBHOOK_URL = None


# ============================================================
# Helpers de negócio
# ============================================================
def email_valido(email: str) -> bool:
    return bool(EMAIL_REGEX.match(email.strip()))


def numero_duimp_valido(numero: str) -> bool:
    return len(re.sub(r"[^0-9A-Za-z]", "", numero)) >= 10


def enviar_para_n8n(url: str, payload: dict) -> requests.Response:
    return requests.post(url, json=payload, timeout=TIMEOUT_REQ,
                         headers={"Content-Type": "application/json"})


def _set_col(ws, row: int, chave: str, valor) -> None:
    """Escreve `valor` na coluna indicada pelo prefixo da chave (ex.: 'C_produto' -> coluna C)."""
    col = chave.split("_", 1)[0]
    if COL_RE.match(col) and valor not in (None, ""):
        ws.cell(row=row, column=column_index_from_string(col), value=valor)


def montar_planilha(planilha: dict) -> bytes:
    """Preenche o template oficial do Omie a partir dos dados estruturados pelo n8n."""
    wb = load_workbook(TEMPLATE_XLSX)
    ws = wb[SHEET_PEDIDO]

    for chave, valor in (planilha.get("cabecalho") or {}).items():
        _set_col(ws, LINHA_CABECALHO, chave, valor)

    linha = LINHA_PRIMEIRO_ITEM
    for item in (planilha.get("itens") or []):
        for chave, valor in item.items():
            _set_col(ws, linha, chave, valor)
        linha += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def extrair_planilha(resp_json):
    """Normaliza a resposta do n8n e retorna o objeto principal."""
    data = resp_json
    if isinstance(data, list):
        data = data[0] if data else {}
    return data if isinstance(data, dict) else {}


# ============================================================
# UI — Header + Hero
# ============================================================
inject(render_template("header", logo_url=resolver_logo_url()))
inject(render_template(
    "hero",
    titulo="Importação de DUIMP",
    subtitulo="Informe o número da DUIMP. O sistema consulta o Portal Único, "
              "prepara os itens no Omie e devolve a planilha de Importação de "
              "Mercadoria pronta para subir.",
))

if not WEBHOOK_URL:
    st.error(
        "⚠️ A URL do webhook N8N não foi configurada. Defina `N8N_WEBHOOK_URL` "
        "em `.streamlit/secrets.toml` ou nos Secrets do Streamlit Cloud."
    )
    st.stop()

if not TEMPLATE_XLSX.exists():
    st.error("⚠️ Template `template_importacao_mercadoria.xlsx` não encontrado na pasta do app.")
    st.stop()


# ============================================================
# UI — Formulário
# ============================================================
email = st.text_input(
    "Email corporativo",
    placeholder=f"usuario@{DOMINIO_PERMITIDO}",
    help=f"Apenas emails @{DOMINIO_PERMITIDO}.",
)
numero_duimp = st.text_input("Número da DUIMP", placeholder="26BR0000407051-9")
versao = st.text_input("Versão", value="1")

st.write("")
gerar = st.button("Gerar planilha", type="primary", use_container_width=True)


# ============================================================
# Lógica
# ============================================================
if gerar:
    erros = []
    if not email.strip():
        erros.append("Informe o email.")
    elif not email_valido(email):
        erros.append(f"Email inválido. Use um endereço @{DOMINIO_PERMITIDO}.")
    if not numero_duimp.strip():
        erros.append("Informe o número da DUIMP.")
    elif not numero_duimp_valido(numero_duimp):
        erros.append("Número da DUIMP inválido.")

    if erros:
        for e in erros:
            st.error(e)
    else:
        with st.spinner("Consultando a DUIMP e preparando a planilha... (pode levar alguns segundos)"):
            try:
                payload = {
                    "email": email.strip(),
                    "numeroDuimp": numero_duimp.strip(),
                    "versao": (versao.strip() or "1"),
                    "tipo": "duimp",
                    "enviado_em": datetime.now().isoformat(timespec="seconds"),
                }
                resp = enviar_para_n8n(WEBHOOK_URL, payload)

                if not (200 <= resp.status_code < 300):
                    st.error(f"O backend respondeu com status {resp.status_code}.")
                    with st.expander("Detalhes da resposta"):
                        st.code(resp.text or "(sem corpo)")
                else:
                    data = extrair_planilha(resp.json())
                    planilha = data.get("planilha")
                    if not planilha or not planilha.get("itens"):
                        st.error("O backend não retornou os itens da planilha.")
                        with st.expander("Resposta recebida"):
                            st.json(data)
                    else:
                        xlsx_bytes = montar_planilha(planilha)
                        n_itens = len(planilha.get("itens", []))
                        num = data.get("numeroDuimp") or numero_duimp.strip()
                        fname = f"Importacao_Mercadoria_DUIMP_{num}.xlsx"

                        st.success(f"Planilha pronta — {n_itens} item(ns) da DUIMP {num}.")
                        if str(data.get("dryRun")).upper() in ("TRUE", "S", "1"):
                            st.info("Modo de teste (dryRun): os produtos podem estar com código provisório.")
                        st.download_button(
                            "⬇️ Baixar planilha de Importação de Mercadoria",
                            data=xlsx_bytes,
                            file_name=fname,
                            mime=MIME_XLSX,
                            use_container_width=True,
                        )
                        with st.expander("Conferir dados retornados"):
                            st.json(planilha)

            except requests.exceptions.Timeout:
                st.error("Tempo de resposta excedido. O fluxo pode estar demorando — tente de novo.")
            except requests.exceptions.ConnectionError:
                st.error("Falha de conexão. Verifique a URL do webhook.")
            except Exception as exc:
                st.error(f"Erro ao montar a planilha: {exc}")


inject(render_template("footer", ano=datetime.now().year))
